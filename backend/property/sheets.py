from __future__ import annotations
import re
import shutil
import threading
from datetime import date, timedelta
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.comments import Comment
from backend.property.config import MONTH_COLUMNS, CATEGORY_ROWS


EXCEL_DATA_DIR = Path(__file__).parent / "data"
EXCEL_TEMPLATE_DIR = Path(__file__).parent / "templates"
# Bundled template is committed to the repo (data/ is gitignored). Desktop
# path is a secondary fallback for fresh dev environments.
EXCEL_TEMPLATE_REPO = EXCEL_TEMPLATE_DIR / "property_manager_2026.xlsx"
EXCEL_TEMPLATE_DESKTOP = Path.home() / "Desktop" / "Property Manager 2026.xlsx"
SOURCE_YEAR = 2026  # Year embodied by the bundled template

# Sheets that are NOT yearly fact tables (carry over as-is on new-year creation,
# data not cleared).
NON_YEARLY_SHEETS = {
    "Property Management Overview",
    "Harga+Biaya Perolehan",
    "- Disclaimer -",
    "Tenants",
    "Reminders",
}


def excel_path_for(year: int) -> Path:
    return EXCEL_DATA_DIR / f"property_manager_{year}.xlsx"


def current_year() -> int:
    return date.today().year


def list_available_years() -> list[int]:
    if not EXCEL_DATA_DIR.exists():
        return []
    years = []
    for p in EXCEL_DATA_DIR.glob("property_manager_*.xlsx"):
        suffix = p.stem.rsplit("_", 1)[-1]
        if suffix.isdigit():
            years.append(int(suffix))
    return sorted(years)


def _to_int(val) -> int:
    """Coerce a cell value to int, returning 0 for None/non-numeric."""
    if val is None:
        return 0
    try:
        return int(float(val))
    except (ValueError, TypeError):
        return 0


class SheetsClient:
    """Local Excel client. One workbook per calendar year.

    Methods accept an optional ``year`` kwarg; when omitted, the current
    calendar year is used. Year-files are auto-created on first access by
    cloning the latest existing year and clearing rental + expense cells.
    Tenants/Reminders sheets carry over (they're cross-year by nature).
    """

    def __init__(self, *_args, **_kwargs):
        EXCEL_DATA_DIR.mkdir(parents=True, exist_ok=True)
        # Re-entrant: writes acquire the lock, then call _load → _ensure_year_file
        # which also acquires it. Without RLock that deadlocks.
        self._lock = threading.RLock()
        self._migrate_legacy_file()

    def _migrate_legacy_file(self) -> None:
        """Rename legacy data/property_manager.xlsx → property_manager_2026.xlsx."""
        legacy = EXCEL_DATA_DIR / "property_manager.xlsx"
        target = excel_path_for(SOURCE_YEAR)
        if legacy.exists() and not target.exists():
            legacy.rename(target)

    def _ensure_year_file(self, year: int) -> Path:
        path = excel_path_for(year)
        if path.exists():
            return path

        with self._lock:
            # Re-check after acquiring lock — another thread may have created it.
            if path.exists():
                return path

            existing = list_available_years()
            if existing:
                template = excel_path_for(existing[-1])
                shutil.copy2(template, path)
                self._clear_year_data(path)
            else:
                template = (
                    EXCEL_TEMPLATE_REPO if EXCEL_TEMPLATE_REPO.exists()
                    else EXCEL_TEMPLATE_DESKTOP if EXCEL_TEMPLATE_DESKTOP.exists()
                    else None
                )
                if template is None:
                    raise FileNotFoundError(
                        f"No template available to create {path.name}. "
                        f"Bundle one at {EXCEL_TEMPLATE_REPO}."
                    )
                shutil.copy2(template, path)
                if year != SOURCE_YEAR:
                    self._clear_year_data(path)
            return path

    @staticmethod
    def _clear_year_data(path: Path) -> None:
        """Wipe all monthly rental income + expense cells from per-property
        sheets while preserving structure, formulas, and global sheets."""
        wb = load_workbook(path)
        try:
            data_rows = list(CATEGORY_ROWS.values())
            month_cols = list(MONTH_COLUMNS.values())

            for sheet_name in wb.sheetnames:
                if sheet_name in NON_YEARLY_SHEETS:
                    continue
                ws = wb[sheet_name]
                for row in data_rows:
                    for col in month_cols:
                        cell = ws[f"{col}{row}"]
                        # Don't clobber formula cells.
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            continue
                        cell.value = None
                        if cell.comment is not None:
                            cell.comment = None
            wb.save(path)
        finally:
            wb.close()

    def _load(self, year: int | None = None, data_only: bool = False):
        y = year if year is not None else current_year()
        path = self._ensure_year_file(y)
        return load_workbook(path, data_only=data_only), y

    def _save(self, wb, year: int) -> None:
        try:
            wb.save(excel_path_for(year))
        finally:
            wb.close()

    @staticmethod
    def _validate_sheet_name(name: str) -> str:
        if not re.match(r'^[\w\s\-]+$', name):
            raise ValueError(f"Invalid sheet name: {name}")
        return name

    @staticmethod
    def _set_income_cell(ws, month: str, amount: int, notes: str) -> None:
        """Set rental-income cell value and replace any prior comment."""
        col = MONTH_COLUMNS[month]
        row = CATEGORY_ROWS["RENTAL INCOME"]
        cell = ws[f"{col}{row}"]
        cell.value = amount
        # Always reset comment so stale annotations don't linger across writes.
        cell.comment = Comment(notes[:255], "PropertyBot") if notes else None

    # ── Writes ────────────────────────────────────────────

    def write_income(self, property_sheet: str, month: str, amount: int,
                     notes: str = "", year: int | None = None) -> None:
        self._validate_sheet_name(property_sheet)
        with self._lock:
            wb, y = self._load(year=year)
            ws = wb[property_sheet]
            self._set_income_cell(ws, month, amount, notes)
            self._save(wb, y)

    def write_expense(self, property_sheet: str, month: str, category: str,
                      amount: int, notes: str = "", year: int | None = None) -> None:
        self._validate_sheet_name(property_sheet)
        with self._lock:
            wb, y = self._load(year=year)
            ws = wb[property_sheet]
            col = MONTH_COLUMNS[month]
            row = CATEGORY_ROWS[category]
            ws[f"{col}{row}"] = amount
            self._save(wb, y)

    def write_income_batch(self, property_sheet: str, months: list[str],
                           amount: int, notes: str = "", year: int | None = None) -> None:
        self._validate_sheet_name(property_sheet)
        with self._lock:
            wb, y = self._load(year=year)
            ws = wb[property_sheet]
            for month in months:
                self._set_income_cell(ws, month, amount, notes)
            self._save(wb, y)

    def write_expense_batch(self, property_sheet: str, months: list[str],
                            amount: int, category: str, year: int | None = None) -> None:
        self._validate_sheet_name(property_sheet)
        with self._lock:
            wb, y = self._load(year=year)
            ws = wb[property_sheet]
            row = CATEGORY_ROWS[category]
            for month in months:
                col = MONTH_COLUMNS[month]
                ws[f"{col}{row}"] = amount
            self._save(wb, y)

    # ── Reads ─────────────────────────────────────────────

    def read_property_summary(self, property_sheet: str, month: str,
                              year: int | None = None) -> dict:
        self._validate_sheet_name(property_sheet)
        with self._lock:
            wb, _ = self._load(year=year, data_only=True)
            try:
                ws = wb[property_sheet]
                col = MONTH_COLUMNS[month]

                def _read(row):
                    return _to_int(ws[f"{col}{row}"].value)

                rental_income = _read(CATEGORY_ROWS["RENTAL INCOME"])
                expenses = {}
                total_expenses = 0
                for cat, row in CATEGORY_ROWS.items():
                    if cat in ("RENTAL INCOME", "OTHER RENTAL INCOME"):
                        continue
                    val = _read(row)
                    if val != 0:
                        expenses[cat] = val
                        total_expenses += val

                return {
                    "property": property_sheet,
                    "month": month,
                    "rental_income": rental_income,
                    "expenses": expenses,
                    "total_expenses": total_expenses,
                    "net": rental_income - total_expenses,
                }
            finally:
                wb.close()

    def read_overview(self, year: int | None = None) -> dict:
        with self._lock:
            wb, _ = self._load(year=year, data_only=True)
            try:
                ws = wb["Property Management Overview"]
                def _r(row, col):
                    return _to_int(ws.cell(row=row, column=col).value)
                return {
                    "ytd_rental_income": _r(5, 15),
                    "ytd_total_expenses": _r(43, 15),
                    "ytd_net": _r(45, 15),
                }
            finally:
                wb.close()

    def get_last_income(self, property_sheet: str, month: str,
                        year: int | None = None) -> int:
        self._validate_sheet_name(property_sheet)
        with self._lock:
            wb, _ = self._load(year=year, data_only=True)
            try:
                ws = wb[property_sheet]
                col = MONTH_COLUMNS[month]
                return _to_int(ws[f"{col}{CATEGORY_ROWS['RENTAL INCOME']}"].value)
            finally:
                wb.close()

    def list_sheets(self, year: int | None = None) -> list[str]:
        with self._lock:
            wb, _ = self._load(year=year, data_only=True)
            try:
                return wb.sheetnames
            finally:
                wb.close()

    # ── Reminders / Tenants (live in current-year file) ──

    def add_reminder(self, reminder, year: int | None = None) -> None:
        with self._lock:
            wb, y = self._load(year=year)
            if "Reminders" not in wb.sheetnames:
                ws = wb.create_sheet("Reminders")
                ws.append(["Property", "Type", "Description", "Due Day", "Amount",
                           "Active", "Last Triggered", "Notes"])
            else:
                ws = wb["Reminders"]
            ws.append([
                reminder.property_key,
                reminder.type,
                reminder.description,
                reminder.due_day,
                reminder.amount,
                "Yes" if reminder.active else "No",
                "",
                reminder.notes,
            ])
            self._save(wb, y)

    def get_active_reminders(self, year: int | None = None) -> list[dict]:
        with self._lock:
            wb, _ = self._load(year=year, data_only=True)
            try:
                if "Reminders" not in wb.sheetnames:
                    return []
                ws = wb["Reminders"]
                headers = [cell.value for cell in ws[1]]
                records = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row[0] is None:
                        continue
                    record = dict(zip(headers, row))
                    if record.get("Active") == "Yes":
                        records.append(record)
                return records
            finally:
                wb.close()

    def add_tenant(self, tenant, year: int | None = None) -> None:
        with self._lock:
            wb, y = self._load(year=year)
            if "Tenants" not in wb.sheetnames:
                ws = wb.create_sheet("Tenants")
                ws.append(["Property", "Name", "Phone", "KTP", "Contract Start",
                           "Contract End", "Monthly Rent", "Deposit",
                           "Deposit Status", "Status", "Notes"])
            else:
                ws = wb["Tenants"]
            ws.append([
                tenant.property_key, tenant.name, tenant.phone, tenant.id_ktp,
                str(tenant.contract_start), str(tenant.contract_end),
                tenant.monthly_rent, tenant.deposit, tenant.deposit_status,
                tenant.status, tenant.notes,
            ])
            self._save(wb, y)

    def get_tenants(self, property_key: str | None = None,
                    year: int | None = None) -> list[dict]:
        with self._lock:
            wb, _ = self._load(year=year, data_only=True)
            try:
                if "Tenants" not in wb.sheetnames:
                    return []
                ws = wb["Tenants"]
                headers = [cell.value for cell in ws[1]]
                records = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row[0] is None:
                        continue
                    record = dict(zip(headers, row))
                    if property_key and record.get("Property") != property_key:
                        continue
                    records.append(record)
                return records
            finally:
                wb.close()

    def get_expiring_leases(self, within_days: int = 30,
                            year: int | None = None) -> list[dict]:
        tenants = self.get_tenants(year=year)
        today = date.today()
        cutoff = today + timedelta(days=within_days)
        expiring = []
        for t in tenants:
            end = t.get("Contract End")
            if not end:
                continue
            if isinstance(end, str):
                try:
                    end = date.fromisoformat(end)
                except ValueError:
                    continue
            if isinstance(end, date) and today <= end <= cutoff:
                t["_days_left"] = (end - today).days
                expiring.append(t)
        return sorted(expiring, key=lambda x: x["_days_left"])
